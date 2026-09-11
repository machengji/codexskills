# -*- coding: utf-8 -*-
"""水利案代码文档生成：只收后端/算法源码，剔除前端样式与页面结构代码。"""
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

PROJECT = Path(r"{{PROJECT}}")
OUT = PROJECT / "{{NAME}}代码.docx"
TXT = Path(r"{{WORKDIR}}\_dam_code_tmp.txt")

# 只收后端与算法：{{PKG}} 下的全部 .py（剔除 tests 的 __init__、__pycache__）
SPECS = [
    ("{{PKG}}/algo", "backend"),
    ("{{PKG}}", "backend"),
]


def strip_code(text: str) -> list[str]:
    text = re.sub(r'"""[\s\S]*?"""', "\n", text)
    text = re.sub(r"'''[\s\S]*?'''", "\n", text)
    lines = []
    for ln in text.splitlines():
        s = ln.rstrip()
        t = s.strip()
        if not t:
            continue
        if t.startswith("#"):
            continue
        lines.append(s)
    return lines


def collect() -> list[tuple[str, list[str]]]:
    files = []
    seen = set()
    for folder, _kind in SPECS:
        base = PROJECT / folder
        if not base.exists():
            continue
        for f in sorted(base.rglob("*.py")):
            p = str(f).replace("\\", "/")
            if "__pycache__" in p:
                continue
            rel = str(f.relative_to(PROJECT)).replace("\\", "/")
            if rel in seen:
                continue
            seen.add(rel)
            lines = strip_code(f.read_text(encoding="utf-8", errors="ignore"))
            if not lines:
                continue
            files.append((rel, lines))
    # 后端在前：算法目录优先，其次服务层，测试放最后
    def rank(item):
        rel = item[0]
        if rel.startswith("{{PKG}}/algo/"):
            return (0, rel)
        if "/tests/" in rel:
            return (2, rel)
        return (1, rel)

    files.sort(key=rank)
    return files


def set_run_font(run, name, size_pt, east=None):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(0, 0, 0)
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), east or name)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size_pt * 2)))
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(int(size_pt * 2)))
    rPr.append(szCs)


def tight_p(p, before=0, after=0, line=240):
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = line / 240.0
    pPr = p._p.get_or_add_pPr()
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))
    spacing.set(qn("w:line"), str(line))
    spacing.set(qn("w:lineRule"), "auto")


def main():
    files, total = collect(), 0
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.left_margin = Mm(19)
    sec.right_margin = Mm(19)
    sec.top_margin = Mm(18)
    sec.bottom_margin = Mm(18)
    for hdr in (sec.header, sec.footer):
        hdr._element.clear()

    blob = []
    for rel, lines in files:
        for s in lines:
            p = doc.add_paragraph()
            tight_p(p, 0, 0, 276)
            run = p.add_run(s[:200])
            set_run_font(run, "Consolas", 12.5, "Consolas")
            blob.append(s)
            total += 1

    TXT.write_text("\n".join(blob), encoding="utf-8")
    doc.save(str(OUT))
    print("files:", len(files))
    print("code lines:", total)
    print("saved:", OUT)


if __name__ == "__main__":
    main()
